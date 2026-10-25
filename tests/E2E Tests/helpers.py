import base64
from pathlib import Path
import openpyxl
from docx import Document


def matching_rows(sheet, text: str) -> list[int]:
    """Row numbers of the cells that equal 'text' exactly."""
    rows = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == text and cell.row not in rows:
                rows.append(cell.row)
    return rows


def get_pan_from_docx(docx_path: Path) -> str | None:
    """Return the next non-empty value after the word 'PAN' in the docx."""
    doc = Document(docx_path)

    texts = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text.strip())
    for paragraph in doc.paragraphs:
        texts.append(paragraph.text.strip())

    for i, text in enumerate(texts):
        if text.strip(": ").upper() == "PAN":
            for next_text in texts[i + 1:]:
                if next_text and next_text.strip(": ").upper() != "PAN":
                    return next_text
    return None


def count_employees_in_excel(excel_file: Path) -> int:
    """One 'Net Salary Paid' row per employee."""
    sheet = openpyxl.load_workbook(excel_file, data_only=True).active
    return len(matching_rows(sheet, "Net Salary Paid"))


def pdfs_in(folder: Path) -> list[Path]:
    return sorted(Path(folder).glob("*.pdf"))


def write_result_files(folder: Path, files: list[dict]) -> None:
    """Save the files returned by generate_payslips (path + base64 data) into a folder."""
    for file in files:
        target = folder / file["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(base64.b64decode(file["data"]))
