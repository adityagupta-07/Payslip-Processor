import pytest
import openpyxl
from openpyxl import Workbook

from src import validate_file_path, load_workbook, get_active_sheet, load_excel_file


def test_validate_file_path_xlsx(tmp_path):
    file_path = tmp_path / "employees.xlsx"
    file_path.touch()
    result = validate_file_path(str(file_path))

    assert result == str(file_path)


def test_validate_file_path_xlsm(tmp_path):
    file_path = tmp_path / "employees.xlsm"
    file_path.touch()
    result = validate_file_path(str(file_path))

    assert result == str(file_path)


def test_validate_file_path_strips_whitespace_and_quotes(tmp_path):
    file_path = tmp_path / "employees.xlsx"
    file_path.touch()
    result = validate_file_path(f'  "{file_path}"  ')

    assert result == str(file_path)


def test_validate_file_path_empty():
    with pytest.raises(ValueError, match="File path cannot be empty."):
        validate_file_path("   ")


def test_validate_file_path_file_does_not_exist(tmp_path):
    file_path = tmp_path / "missing.xlsx"
    with pytest.raises(FileNotFoundError, match="File does not exist."):
        validate_file_path(str(file_path))


def test_validate_file_path_path_is_directory(tmp_path):
    directory = tmp_path / "employees.xlsx"
    directory.mkdir()
    with pytest.raises(ValueError, match="Path is not a file."):
        validate_file_path(str(directory))


def test_validate_file_path_invalid_extension(tmp_path):
    file_path = tmp_path / "employees.txt"
    file_path.touch()
    with pytest.raises(
        ValueError,
        match="Please provide an .xlsx or .xlsm Excel file.",
    ):
        validate_file_path(str(file_path))


def test_load_workbook(tmp_path):
    file_path = tmp_path / "employees.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "John Doe"
    workbook.save(file_path)
    result = load_workbook(str(file_path))

    assert result.active["A1"].value == "John Doe"


def test_load_workbook_missing_file(tmp_path):
    file_path = tmp_path / "missing.xlsx"
    with pytest.raises(FileNotFoundError):
        load_workbook(str(file_path))


def test_get_active_sheet():
    workbook = Workbook()
    workbook.active.title = "Employees"
    result = get_active_sheet(workbook)

    assert result.title == "Employees"


def test_get_active_sheet_returns_active_sheet():
    workbook = Workbook()
    workbook.active.title = "Employees"
    workbook.create_sheet("Other")
    workbook.active = 1
    result = get_active_sheet(workbook)

    assert result.title == "Other"


def test_load_excel_file(tmp_path):
    file_path = tmp_path / "employees.xlsx"
    workbook = Workbook()
    workbook.active.title = "Employees"
    workbook.active["A1"] = "John Doe"
    workbook.save(file_path)
    result = load_excel_file(str(file_path))

    assert result["A1"].value == "John Doe"


def test_load_excel_file_missing_file(tmp_path):
    file_path = tmp_path / "missing.xlsx"
    with pytest.raises(FileNotFoundError):
        load_excel_file(str(file_path))
