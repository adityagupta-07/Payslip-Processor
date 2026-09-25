from .config import Employee
from .constants import ExcelConstants
from typing import Optional, Union
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

excel_constants = ExcelConstants()


def matching_row_numbers(search_string: str, sheet: Worksheet) -> list[int]:
    matching_rows: list[int] = []
    for target_row in sheet.iter_rows():
        for cell in target_row:
            if cell.value == search_string:
                if cell.row not in matching_rows:
                    matching_rows.append(cell.row)
    return matching_rows


def get_employee_block(sheet: Worksheet, string1: str, string2: str) -> list[tuple[int, int]]:
    string1_rows = matching_row_numbers(string1, sheet)
    string2_rows = matching_row_numbers(string2, sheet)
    return list(zip(string1_rows, string2_rows))


def normalize_cell_value(value: object) -> object:
    return value.strip() if isinstance(value, str) else value


def set_ssf_employer_value(
    employee_obj: Employee,
    sheet: Worksheet,
    r: int,
    target_cell: Cell,
    next_cell: Cell,
    already_found: bool,
) -> None:
    if already_found:
        employee_obj.set_atr(
            excel_constants.get_atr("SSF_EMPLOYER_LOWER"), f"Rs. {next_cell.value:,.2f}"
        )
    else:
        employee_obj.set_atr(
            excel_constants.get_atr("SSF_EMPLOYER_UPPER"), f"Rs. {next_cell.value:,.2f}"
        )
        employee_obj.set_atr(
            excel_constants.get_atr("FINANCIAL_YEAR_NOTE"),
            sheet.cell(row=r, column=(target_cell.column + 2)).value,
        )


def set_numeric_or_plain_value(
    employee_obj: Employee,
    target_value: object,
    next_cell: Cell,
    r: int,
    block: tuple[int, int],
) -> None:
    if isinstance(next_cell.value, (int, float)) and (r > block[1] - 13 and r < block[1] + 1):
        employee_obj.set_atr(target_value, f"Rs. {next_cell.value:,.2f}")
    else:
        employee_obj.set_atr(target_value, next_cell.value)


def set_month_year(
    employee_obj: Employee,
    r: int,
    target_cell: Cell,
    target_value: object,
    next_cell: Cell,
    block: tuple[int, int],
) -> None:
    if r == (block[1] - 2) and target_cell.column == 3:
        employee_obj.set_atr(excel_constants.get_atr("MONTH_YEAR"), next_cell.value)
        employee_obj.set_atr(target_value, next_cell.value.strftime("%b %Y"))


def get_details_per_employee(
    sheet: Worksheet, block: tuple[int, int], employee_obj: Employee
) -> Employee:

    found_ssf_contribution_by_employer = False

    for r in range(block[0], block[1] + 1):
        for target_cell in sheet[r]:

            target_value = normalize_cell_value(target_cell.value)

            if not employee_obj.check_attribute(str(target_value)):
                continue

            next_cell = sheet.cell(row=r, column=(target_cell.column + 1))

            if excel_constants.get_atr("SSF_EMPLOYER") in target_value:
                set_ssf_employer_value(
                    employee_obj, sheet, r, target_cell, next_cell, found_ssf_contribution_by_employer
                )
                if found_ssf_contribution_by_employer:
                    break
                found_ssf_contribution_by_employer = True

            set_numeric_or_plain_value(employee_obj, target_value, next_cell, r, block)

            set_month_year(employee_obj, r, target_cell, target_value, next_cell, block)

    return employee_obj