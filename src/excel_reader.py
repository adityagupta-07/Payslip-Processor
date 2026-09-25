import openpyxl
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def validate_file_path(path: str) -> str:
    """Validate and return a valid Excel file path."""

    path = path.strip().replace('"', "")

    if not path:
        raise ValueError("File path cannot be empty.")

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError("File does not exist.")

    if not file_path.is_file():
        raise ValueError("Path is not a file.")

    if file_path.suffix.lower() not in [".xlsx", ".xlsm"]:
        raise ValueError("Please provide an .xlsx or .xlsm Excel file.")    
    return str(file_path)


def user_input() -> str:
    """Prompt the user until they provide a valid Excel file path."""

    while True:
        try:
            path = input("Provide file path: ")
            return validate_file_path(path)
        except (ValueError, FileNotFoundError) as error:
            print(f"Error: {error}")


def load_workbook(file_path: str) -> openpyxl.Workbook:
    """Load an Excel workbook with cell values instead of formulas."""
    return openpyxl.load_workbook(file_path, data_only=True)


def get_active_sheet(workbook: openpyxl.Workbook) -> Worksheet:
    """Return the active worksheet from the workbook."""
    return workbook.active


def load_excel_file(file_path: str) -> Worksheet:
    """Load the Excel file and return its active worksheet."""
    workbook = load_workbook(file_path)
    return get_active_sheet(workbook)